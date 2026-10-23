"""Unit tests for Excel Report Generation & Yellow Cell Highlighting."""

import io
import unittest
import openpyxl
from core.excel_generator import ExcelReportGenerator


class TestExcelReportGenerator(unittest.TestCase):

    def setUp(self):
        self.profile = {
            "company_type": "Public Company",
            "listing_status": "Listed",
            "framework": "Indian Accounting Standards (Ind AS)",
            "schedule_iii_division": "Division II",
            "financial_year": "FY 2023-24",
            "statement_type": "Standalone",
        }
        self.mock_issues = [
            {
                "issue_id": "FRR-A-001",
                "category": "Arithmetic & Balance Sheet Integrity",
                "ref": "Balance Sheet",
                "description": "Balance Sheet Equation Out of Balance: Difference of 250.00 detected.",
                "current_disclosure": "Assets = 11,060 vs Liab = 10,810",
                "suggested_correction": "Reconcile Assets and Liabilities.",
                "statutory_ref": "Section 129(1) of Companies Act, 2013",
                "action": "Immediate Pre-Finalisation Correction Required.",
                "status": "Critical Correction",
                "risk": "High",
                "highlight_cells": ["Total Assets", "Total Equity and Liabilities"]
            },
            {
                "issue_id": "FRR-A-002",
                "category": "Schedule III - Mandatory Ageing Schedules",
                "ref": "Note 14 / Trade Payables",
                "description": "Mandatory Trade Payables Ageing Schedule is missing.",
                "current_disclosure": "Presented without ageing.",
                "suggested_correction": "Insert MCA mandatory ageing table.",
                "statutory_ref": "Schedule III, Division II, Part I",
                "action": "Draft and include mandatory Trade Payables ageing table.",
                "status": "Mandatory Disclosure Missing",
                "risk": "High",
                "highlight_cells": ["Trade Payables Note - Ageing Schedule"]
            }
        ]

    def test_option_a_workbook_structure_and_yellow_highlight(self):
        """Option A workbook must have 8 columns in Issue Log and yellow highlighted cells in Sheet 2."""
        excel_bytes = ExcelReportGenerator.generate_option_a_workbook(self.mock_issues, self.profile)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))

        # Check Sheet 1
        ws1 = wb.worksheets[0]
        self.assertEqual(ws1.title, "Issue Log (Pre-Finalisation)")

        expected_columns = [
            "Issue ID",
            "Financial statement or note reference",
            "Description of the issue",
            "Current reference or existing disclosure",
            "Correct reference or suggested correction",
            "Applicable accounting standard, Schedule III requirement or statutory provision",
            "Recommended action",
            "Status of the issue",
        ]

        header_row = 4
        for col_idx, expected_col in enumerate(expected_columns, start=1):
            actual_val = ws1.cell(row=header_row, column=col_idx).value
            self.assertEqual(actual_val, expected_col, f"Column {col_idx} header does not match!")

        # Verify yellow fill on suggested correction (Col 5)
        cell_c5 = ws1.cell(row=5, column=5)
        self.assertIsNotNone(cell_c5.fill.start_color.rgb)

        # Check Sheet 2
        self.assertGreaterEqual(len(wb.worksheets), 2)
        ws2 = wb.worksheets[1]
        self.assertEqual(ws2.title, "Financial Statements (Marked)")

        # Verify presence of yellow highlighted cells in Sheet 2
        yellow_cell_found = False
        for row in ws2.iter_rows(min_row=5, max_row=40):
            for cell in row:
                if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb:
                    color_rgb = str(cell.fill.start_color.rgb)
                    # Check for yellow: FFFF00 or FFF2CC
                    if "FFFF00" in color_rgb or "FFF2CC" in color_rgb:
                        yellow_cell_found = True
                        break
            if yellow_cell_found:
                break

        self.assertTrue(yellow_cell_found, "Worksheet 2 does not contain required yellow highlighted cells!")

    def test_option_b_workbook_structure(self):
        """Option B workbook must have Red-Flag Register and Executive Summary."""
        stats = {
            "total_issues": 2,
            "high_risk_count": 2,
            "med_risk_count": 0,
            "low_risk_count": 0,
            "categories_breakdown": {"Arithmetic": 1, "Schedule III": 1}
        }
        excel_bytes = ExcelReportGenerator.generate_option_b_workbook(self.mock_issues, self.profile, stats)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))

        # Check Sheet 1
        ws1 = wb.worksheets[0]
        self.assertEqual(ws1.title, "Red-Flag & Findings Register")
        self.assertEqual(ws1.cell(row=4, column=1).value, "Issue ID")
        self.assertEqual(ws1.cell(row=4, column=7).value, "Risk Assessment")

        # Check Sheet 2
        ws2 = wb.worksheets[1]
        self.assertEqual(ws2.title, "Executive Summary & Analytics")
        self.assertIn("EXECUTIVE DASHBOARD", ws2["A1"].value)

    def test_review_points_workbook_structure(self):
        """Review Points workbook must have 8 columns in Sheet 1 and Summary in Sheet 2."""
        excel_bytes = ExcelReportGenerator.generate_review_points_workbook(
            issues=self.mock_issues,
            profile=self.profile
        )
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertEqual(len(wb.worksheets), 2)
        ws1 = wb.worksheets[0]
        self.assertEqual(ws1.title, "Review Points Sheet")
        self.assertEqual(ws1.cell(row=4, column=1).value, "Issue ID")
        self.assertEqual(ws1.cell(row=4, column=8).value, "Status of the issue")
        ws2 = wb.worksheets[1]
        self.assertEqual(ws2.title, "Executive Summary & Analytics")

    def test_schedule_iii_draft_fs_workbook_structure_and_yellow_highlight(self):
        """Schedule III Draft FS workbook must have 4 sheets and yellow highlighted additions/modifications."""
        excel_bytes = ExcelReportGenerator.generate_schedule_iii_draft_fs_workbook(
            issues=self.mock_issues,
            profile=self.profile
        )
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        expected_titles = [
            "Balance Sheet (Schedule III)",
            "Statement of Profit & Loss",
            "Notes & Ageing Schedules",
            "Summary of Modifications"
        ]
        self.assertEqual([ws.title for ws in wb.worksheets], expected_titles)

        # Check yellow highlight in Balance Sheet
        ws_bs = wb["Balance Sheet (Schedule III)"]
        bs_yellow = False
        for row in ws_bs.iter_rows(min_row=5, max_row=45):
            for cell in row:
                if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb:
                    color = str(cell.fill.start_color.rgb)
                    if "FFFF00" in color or "FFF2CC" in color:
                        bs_yellow = True
                        break
            if bs_yellow:
                break
        self.assertTrue(bs_yellow, "Balance Sheet does not contain yellow highlighted additions/modifications!")

        # Check yellow highlight in Notes & Ageing Schedules
        ws_notes = wb["Notes & Ageing Schedules"]
        notes_yellow = False
        for row in ws_notes.iter_rows(min_row=5, max_row=60):
            for cell in row:
                if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb:
                    color = str(cell.fill.start_color.rgb)
                    if "FFFF00" in color or "FFF2CC" in color:
                        notes_yellow = True
                        break
            if notes_yellow:
                break
        self.assertTrue(notes_yellow, "Notes & Ageing Schedules does not contain yellow highlights!")


if __name__ == "__main__":
    unittest.main()
